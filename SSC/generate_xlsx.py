import shutil
from openpyxl import load_workbook
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = ROOT / 'SSC' / 'digital_io.xlsx'
BAK_PATH = ROOT / 'SSC' / 'digital_io_bak.xlsx'

def backup_xlsx():
    if XLSX_PATH.exists() and not BAK_PATH.exists():
        shutil.copy2(XLSX_PATH, BAK_PATH)

def find_row_by_index_prefix(ws, prefix):
    for row in range(1, ws.max_row + 1):
        val = ws.cell(row, 2).value
        # sometimes val is an int or float if excel parsed it weirdly, but usually hex string
        if val and isinstance(val, str) and val.strip().lower().startswith(prefix.lower()):
            return row
    return None

def clear_and_insert_dict_object(ws, insert_at_row, index_val, name, access, rxtx, obj_code='RECORD'):
    ws.insert_rows(insert_at_row, 10)
    
    curr_row = insert_at_row
    ws.cell(curr_row, 2, index_val)
    ws.cell(curr_row, 3, obj_code)
    ws.cell(curr_row, 6, name)
    ws.cell(curr_row, 12, access)
    ws.cell(curr_row, 13, rxtx)
    ws.cell(curr_row, 10, 'M') # Auto Mapping
    curr_row += 1
    
    sub_items = [
        ('1', 'UINT', 'CanIdLo', '0'),
        ('2', 'UINT', 'CanIdHi', '0'),
        ('3', 'USINT', 'Dlc', '0'),
        ('4', 'USINT', 'Flags', '0'),
        ('5', 'UINT', 'Reserved', '0'),
        ('6', 'STRING(30)', 'Payload1', ''),
        ('7', 'STRING(30)', 'Payload2', ''),
        ('8', 'STRING(4)', 'Payload3', ''),
    ]
    for si, dtype, si_name, default in sub_items:
        ws.cell(curr_row, 4, si)
        ws.cell(curr_row, 5, dtype)
        ws.cell(curr_row, 6, si_name)
        ws.cell(curr_row, 7, default)
        ws.cell(curr_row, 12, access)
        curr_row += 1
        
    return insert_at_row + 10 

def insert_config_object(ws, insert_at_row, index_val, name):
    ws.insert_rows(insert_at_row, 7)
    curr_row = insert_at_row
    ws.cell(curr_row, 2, index_val)
    ws.cell(curr_row, 3, 'RECORD')
    ws.cell(curr_row, 6, name)
    curr_row += 1
    
    sub_items = [
        ('1', 'USINT', 'Enable', '0'),
        ('2', 'USINT', 'FdEnable', '1'),
        ('3', 'UDINT', 'BitrateArb', '2000000'),
        ('4', 'UDINT', 'BitrateData', '8000000'),
        ('5', 'USINT', 'CleanDirtyFrame', '0'),
    ]
    for si, dtype, si_name, default in sub_items:
        ws.cell(curr_row, 4, si)
        ws.cell(curr_row, 5, dtype)
        ws.cell(curr_row, 6, si_name)
        ws.cell(curr_row, 7, default)
        ws.cell(curr_row, 12, 'rw') # rw per spec
        curr_row += 1
    return insert_at_row + 7

def insert_status_object(ws, insert_at_row, index_val, name):
    ws.insert_rows(insert_at_row, 6)
    curr_row = insert_at_row
    ws.cell(curr_row, 2, index_val)
    ws.cell(curr_row, 3, 'RECORD')
    ws.cell(curr_row, 6, name)
    curr_row += 1
    
    sub_items = [
        ('1', 'USINT', 'BusState', '0'),
        ('2', 'USINT', 'ErrorMode', '0'),
        ('3', 'UINT', 'Tec', '0'),
        ('4', 'UINT', 'Rec', '0'),
    ]
    for si, dtype, si_name, default in sub_items:
        ws.cell(curr_row, 4, si)
        ws.cell(curr_row, 5, dtype)
        ws.cell(curr_row, 6, si_name)
        ws.cell(curr_row, 7, default)
        ws.cell(curr_row, 12, 'ro') # ro per spec
        curr_row += 1
    return insert_at_row + 6

def run_gen():
    backup_xlsx()
    wb = load_workbook(BAK_PATH)
    
    # 1. Enable FOE
    for sheet_name in ['Options', 'EtherCAT', 'Profile']:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for r in range(1, 100):
                if ws.cell(r, 2).value == 'FOE_SUPPORTED':
                    ws.cell(r, 3, 1)
                    break
    
    ws_prof = wb['Profile']
    
    # --- 0x6000 - 0x6030 ---
    row_60 = find_row_by_index_prefix(ws_prof, '0x6000')
    if row_60: ws_prof.delete_rows(row_60, 10)
    else: row_60 = 35
    curr = row_60
    for i in range(4):
        idx = f'0x{0x6000 + i*0x10:04X}'
        curr = clear_and_insert_dict_object(ws_prof, curr, idx, f'Can{i}RxFrame', 'ro', 'tx')
        
    # --- 0x7000 - 0x7030 ---
    row_70 = find_row_by_index_prefix(ws_prof, '0x7000')
    if row_70: ws_prof.delete_rows(row_70, 10)
    else: row_70 = curr + 5
    curr = row_70
    for i in range(4):
        idx = f'0x{0x7000 + i*0x10:04X}'
        curr = clear_and_insert_dict_object(ws_prof, curr, idx, f'Can{i}TxFrame', 'rw', 'rx')

    # --- 0x8000 - 0x8030 ---
    row_80 = find_row_by_index_prefix(ws_prof, '0x8000')
    if row_80: ws_prof.delete_rows(row_80, 6) # approx previous height ~6
    else: row_80 = curr + 5
    curr = row_80
    for i in range(4):
        idx = f'0x{0x8000 + i*0x10:04X}'
        curr = insert_config_object(ws_prof, curr, idx, f'Can{i}Config')

    # --- 0x9000 - 0x9030 ---
    row_90 = find_row_by_index_prefix(ws_prof, '0x9000')
    if row_90: ws_prof.delete_rows(row_90, 6) # approx previous height ~6
    else: row_90 = curr + 5
    curr = row_90
    for i in range(4):
        idx = f'0x{0x9000 + i*0x10:04X}'
        curr = insert_status_object(ws_prof, curr, idx, f'Can{i}Status')
        
    wb.save(XLSX_PATH)
    print('XLSX update complete. SDOs config/status are now extended to 4 channels.')

if __name__ == '__main__':
    run_gen()

